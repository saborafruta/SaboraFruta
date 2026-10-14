"""Exportações da DRE gerencial usando exatamente os dados exibidos na tela."""
from __future__ import annotations

import io
from html import escape

from django.http import HttpResponse


ROTULOS_REGIME = {
    'caixa': 'Caixa (data do pagamento)',
    'competencia': 'Competência',
    'vencimento': 'Vencimento',
}


def _numero(valor):
    return float(valor or 0)


def _moeda(valor):
    texto = f'{_numero(valor):,.2f}'
    return 'R$ ' + texto.replace(',', '#').replace('.', ',').replace('#', '.')


def _linhas(dados):
    """Linhas reconciliadas, compartilhadas pelo Excel e pelo PDF."""
    for bloco in dados['linhas']:
        yield {'tipo': 'bloco', 'nome': bloco['rotulo']}
        for linha in bloco['linhas']:
            yield {
                'tipo': 'grupo',
                'nome': f"{linha['codigo']} · {linha['descricao']}".strip(' ·'),
                'mes': linha['mes'],
                'participacao': linha['participacao_pct'],
                'anterior': linha['anterior'],
                'variacao': linha['variacao'],
                'ano': linha['ano'],
            }
            for filha in linha['filhas']:
                yield {
                    'tipo': 'filha',
                    'nivel': filha.get('nivel', 1),
                    'nome': f"{filha['codigo']} · {filha['descricao']}".strip(' ·'),
                    'mes': filha['total'],
                    'participacao': filha['participacao_pct'],
                    'anterior': None,
                    'variacao': None,
                    'ano': None,
                }
        sem = bloco['sem_categoria']
        if sem['tem']:
            yield {
                'tipo': 'sem_categoria', 'nome': 'Sem categoria',
                'mes': sem['mes'], 'participacao': sem['participacao_pct'],
                'anterior': sem['anterior'], 'variacao': None, 'ano': sem['ano'],
            }
        yield {
            'tipo': 'total', 'nome': f"Total de {bloco['rotulo'].lower()}",
            'mes': bloco['total_mes'],
            'participacao': 100 if bloco['total_mes'] else None,
            'anterior': bloco['total_anterior'], 'variacao': None,
            'ano': bloco['total_ano'],
        }

    resumo = dados['resumo']
    yield {
        'tipo': 'resultado', 'nome': 'Resultado do período',
        'mes': resumo['resultado'], 'participacao': resumo['margem'],
        'anterior': resumo['resultado_anterior'],
        'variacao': resumo['variacao_resultado'], 'ano': resumo['resultado_ano'],
    }


def exportar_excel(dados, filial):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    competencia = dados['competencia']
    wb = Workbook()
    ws = wb.active
    ws.title = 'DRE'
    ws.sheet_view.showGridLines = False

    empresa = getattr(filial, 'razao_social', '') or str(filial)
    ws.append(['DRE Gerencial'])
    ws.append([empresa])
    ws.append([
        f"Período: {competencia.strftime('%m/%Y')}",
        f"Critério: {ROTULOS_REGIME.get(dados['regime'], dados['regime'])}",
    ])
    ws.append([
        f"Acumulado: janeiro a {competencia.strftime('%m/%Y')}",
    ])
    ws.append([])
    ws.append([
        'Grupo e categoria', competencia.strftime('%m/%Y'), 'Participação',
        'Mês anterior', 'Variação', 'Acumulado no ano',
    ])

    azul = PatternFill('solid', fgColor='1D4ED8')
    azul_claro = PatternFill('solid', fgColor='DBEAFE')
    cinza = PatternFill('solid', fgColor='E2E8F0')
    vermelho_claro = PatternFill('solid', fgColor='FEE2E2')
    for cell in ws[1]:
        cell.font = Font(size=16, bold=True, color='1D4ED8')
    for cell in ws[6]:
        cell.fill = azul
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center')

    for item in _linhas(dados):
        if item['tipo'] == 'bloco':
            ws.append([item['nome']])
        else:
            nome = ('   ' * item.get('nivel', 0)) + item['nome']
            ws.append([
                nome, _numero(item['mes']),
                (_numero(item['participacao']) / 100
                 if item['participacao'] is not None else None),
                _numero(item['anterior']) if item['anterior'] is not None else None,
                _numero(item['variacao']) if item['variacao'] is not None else None,
                _numero(item['ano']) if item['ano'] is not None else None,
            ])
        linha_excel = ws.max_row
        if item['tipo'] == 'bloco':
            for cell in ws[linha_excel]:
                cell.fill = azul_claro
                cell.font = Font(bold=True, color='1E3A8A')
        elif item['tipo'] in {'total', 'resultado'}:
            for cell in ws[linha_excel]:
                cell.fill = vermelho_claro if item['tipo'] == 'resultado' else cinza
                cell.font = Font(bold=True)
        elif item['tipo'] == 'grupo':
            ws.cell(linha_excel, 1).font = Font(bold=True)

    for row in range(7, ws.max_row + 1):
        for col in (2, 4, 5, 6):
            ws.cell(row, col).number_format = 'R$ #,##0.00;[Red]-R$ #,##0.00'
        ws.cell(row, 3).number_format = '0.0%'
    ws.freeze_panes = 'B7'
    larguras = (48, 16, 15, 16, 16, 20)
    for indice, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = largura

    conteudo = io.BytesIO()
    wb.save(conteudo)
    resposta = HttpResponse(
        conteudo.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    resposta['Content-Disposition'] = (
        f'attachment; filename="dre-{competencia:%Y-%m}-{dados["regime"]}.xlsx"'
    )
    return resposta


def exportar_pdf(dados, filial):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    competencia = dados['competencia']
    conteudo = io.BytesIO()
    documento = SimpleDocTemplate(
        conteudo, pagesize=landscape(A4),
        rightMargin=12 * mm, leftMargin=12 * mm,
        topMargin=12 * mm, bottomMargin=12 * mm,
        title=f'DRE Gerencial {competencia:%m/%Y}',
    )
    estilos = getSampleStyleSheet()
    pequeno = ParagraphStyle(
        'dre-pequeno', parent=estilos['BodyText'], fontSize=7.5,
        leading=9, textColor=colors.HexColor('#334155'),
    )
    titulo_linha = ParagraphStyle(
        'dre-titulo-linha', parent=pequeno, fontName='Helvetica-Bold',
    )
    empresa = getattr(filial, 'razao_social', '') or str(filial)
    elementos = [
        Paragraph('DRE Gerencial', estilos['Title']),
        Paragraph(escape(empresa), estilos['Heading3']),
        Paragraph(
            f"Período: <b>{competencia:%m/%Y}</b> &nbsp;&nbsp; "
            f"Critério: <b>{escape(ROTULOS_REGIME.get(dados['regime'], dados['regime']))}</b> &nbsp;&nbsp; "
            f"Acumulado: <b>janeiro a {competencia:%m/%Y}</b>",
            estilos['BodyText'],
        ),
        Spacer(1, 5 * mm),
    ]

    tabela = [[
        'Grupo e categoria', competencia.strftime('%m/%Y'), 'Participação',
        'Mês anterior', 'Variação', 'Acumulado no ano',
    ]]
    tipos = []
    for item in _linhas(dados):
        tipos.append(item['tipo'])
        if item['tipo'] == 'bloco':
            tabela.append([Paragraph(escape(item['nome']), titulo_linha), '', '', '', '', ''])
            continue
        recuo = '&nbsp;' * (item.get('nivel', 0) * 5)
        tabela.append([
            Paragraph(recuo + escape(item['nome']),
                      titulo_linha if item['tipo'] in {'grupo', 'total', 'resultado'} else pequeno),
            _moeda(item['mes']),
            f"{item['participacao']}%" if item['participacao'] is not None else '—',
            _moeda(item['anterior']) if item['anterior'] is not None else '',
            _moeda(item['variacao']) if item['variacao'] is not None else '',
            _moeda(item['ano']) if item['ano'] is not None else '',
        ])

    grade = Table(
        tabela, repeatRows=1,
        colWidths=[105 * mm, 30 * mm, 25 * mm, 30 * mm, 30 * mm, 35 * mm],
    )
    estilo = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1D4ED8')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 7.5),
        ('ALIGN', (1, 0), (-1, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]
    for indice, tipo in enumerate(tipos, start=1):
        if tipo == 'bloco':
            estilo.extend([
                ('SPAN', (0, indice), (-1, indice)),
                ('BACKGROUND', (0, indice), (-1, indice), colors.HexColor('#DBEAFE')),
                ('TEXTCOLOR', (0, indice), (-1, indice), colors.HexColor('#1E3A8A')),
            ])
        elif tipo in {'total', 'resultado'}:
            estilo.extend([
                ('FONTNAME', (0, indice), (-1, indice), 'Helvetica-Bold'),
                ('BACKGROUND', (0, indice), (-1, indice),
                 colors.HexColor('#FEE2E2' if tipo == 'resultado' else '#E2E8F0')),
            ])
    grade.setStyle(TableStyle(estilo))
    elementos.append(grade)
    documento.build(elementos)

    resposta = HttpResponse(conteudo.getvalue(), content_type='application/pdf')
    resposta['Content-Disposition'] = (
        f'attachment; filename="dre-{competencia:%Y-%m}-{dados["regime"]}.pdf"'
    )
    return resposta
